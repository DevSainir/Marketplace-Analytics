"""The P&L export must carry the right columns and re-import cleanly.

Written with ``xlsxwriter``, read back here with ``openpyxl`` (a different
library on purpose — a bug shared by writer and reader could not hide a broken
column). Two properties carry the weight: the sheet shows the "By SKU"
columns the screen shows, and a freshly-exported file re-imports — the editable
"Cost price per unit" column round-trips, while the "TOTAL"/bucket summary
rows are keyed blank so a re-upload skips them instead of treating 0 as a SKU.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from src.modules.finreports.schemas.sku_report import (
    PnlDashboard,
    PnlSkuReportResponse,
    PnlSkuRow,
)
from src.modules.finreports.services.cost_price_workbook import (
    COST_HEADER,
    NM_ID_HEADER,
    build_pnl_workbook,
)
from src.modules.finreports.services.cost_upload_parser import parse_cost_upload


def _row(nm_id: int, sa_name: str, *, cost: str | None) -> PnlSkuRow:
    """A per-SKU P&L row with every field populated (cost optional)."""
    cost_dec = Decimal(cost) if cost is not None else None
    return PnlSkuRow(
        nm_id=nm_id,
        sa_name=sa_name,
        subject_name="T-shirt",
        brand_name="ACME",
        is_excluded=False,
        orders_units=10,
        sold_units=8,
        returned_units=1,
        return_rate_pct=Decimal("12.5"),
        buyout_units=8,
        decided_units=20,
        buyout_pct=Decimal("40"),
        sales_amount=Decimal("1000"),
        returns_amount=Decimal("100"),
        revenue_paid=Decimal("900"),
        revenue_seller_price=Decimal("1100"),
        spp_amount=Decimal("200"),
        spp_pct=Decimal("18"),
        average_cheque=Decimal("112.5"),
        logistics_forward=Decimal("50"),
        logistics_return=Decimal("10"),
        logistics_total=Decimal("60"),
        logistics_pct=Decimal("6.6"),
        penalties_total=Decimal("0"),
        penalties=[],
        cost_price=cost_dec,
        cogs=(cost_dec * 8 if cost_dec is not None else None),
        advertising=Decimal("30"),
        advertising_pct=Decimal("3.3"),
        ad_bonus=Decimal("12"),
        advertising_cash=Decimal("18"),
        payout_for_goods=Decimal("840"),
        payout_net=Decimal("780"),
        net_profit=(Decimal("100") if cost_dec is not None else None),
        net_profit_pct=None,
        roi_pct=None,
    )


def _report(rows: list[PnlSkuRow], *, unallocated: PnlSkuRow | None = None) -> PnlSkuReportResponse:
    """Wrap rows into a report with a totals row and a dashboard."""
    dash = PnlDashboard(
        sales=Decimal("1000"),
        paid=Decimal("900"),
        sold_units=8,
        returned_units=1,
        return_rate_pct=Decimal("12.5"),
        logistics=Decimal("60"),
        logistics_pct=Decimal("6.6"),
        cogs=Decimal("5600"),
        uncosted_count=1,
        payout_net=Decimal("780"),
        other_deductions=Decimal("30"),
        advertising=Decimal("30"),
        excluded_count=0,
    )
    return PnlSkuReportResponse(
        date_from=date(2024, 1, 1),
        date_to=date(2026, 1, 1),
        currency="KGS",
        rows=rows,
        unallocated=unallocated,
        totals=_row(0, "Total", cost="700"),
        dashboard=dash,
        notes=[],
    )


def _header(data: bytes) -> list[str]:
    """The first row of the exported sheet, read with the other library."""
    ws = load_workbook(io.BytesIO(data), read_only=True).active
    assert ws is not None
    return [cell.value for cell in next(ws.iter_rows())]


def test_the_export_carries_the_pnl_columns() -> None:
    # Arrange & Act.
    data = build_pnl_workbook(_report([_row(111, "w-lsx01-dbl", cost="700")]))

    # Assert: the "By SKU" identity + margin columns are present and named.
    header = _header(data)
    for expected in (
        "Vendor code",
        NM_ID_HEADER,
        "Category",
        "Brand",
        "Sales",
        "Logistics",
        COST_HEADER,
        "Payout",
        "Net profit",
    ):
        assert expected in header
    # And the two columns the redesign dropped.
    assert "Currency" not in header
    assert "Source" not in header


def test_a_costed_and_an_uncosted_sku_round_trip() -> None:
    # Arrange.
    report = _report([_row(111, "w-lsx01-dbl", cost=None), _row(222, "m-lnx02-blk", cost="662.50")])

    # Act: export, then read the bytes back through the real parser.
    parsed = parse_cost_upload("pnl.xlsx", build_pnl_workbook(report))

    # Assert: both SKUs come back; cost preserved, gap blank. The "TOTAL" row is
    # keyed blank, so it is not among the parsed rows.
    by_nm = {int(p.nm_id_raw): p for p in parsed}
    assert set(by_nm) == {111, 222}
    assert by_nm[111].cost_raw is None
    assert Decimal(str(by_nm[222].cost_raw)) == Decimal("662.50")


def test_the_bucket_and_totals_rows_are_not_parsed_as_skus() -> None:
    # Arrange: a report that carries an account-level bucket and a totals row.
    report = _report(
        [_row(111, "w-lsx01-dbl", cost="700")],
        unallocated=_row(0, "", cost="500"),
    )

    # Act.
    parsed = parse_cost_upload("pnl.xlsx", build_pnl_workbook(report))

    # Assert: only the real SKU — the bucket (nm_id 0) and "TOTAL" are blank-keyed.
    assert [int(p.nm_id_raw) for p in parsed] == [111]
